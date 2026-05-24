from __future__ import annotations

import re
from abc import ABC, abstractmethod
from pathlib import Path
from statistics import median
from typing import Any

import fitz
from docx import Document

from models import DocumentMetadata, ExtractedDocument, ExtractedSection
from utils import get_logger


logger = get_logger(__name__)


class DocumentExtractor(ABC):
    @abstractmethod
    def extract(self, file_path: str | Path, document_name: str | None = None) -> ExtractedDocument:
        raise NotImplementedError


class TxtExtractor(DocumentExtractor):
    def extract(self, file_path: str | Path, document_name: str | None = None) -> ExtractedDocument:
        path = Path(file_path)
        filename = document_name or path.name
        lines = path.read_text(encoding="utf-8").splitlines()
        sections = _sections_from_text_lines(
            lines=lines,
            document_name=filename,
            page_number=1,
            metadata={"extraction_backend": "text"},
        )
        metadata = DocumentMetadata(filename=filename, file_type=".txt", total_pages=1)
        logger.info("Extracted TXT filename=%s sections=%s", filename, len(sections))
        return ExtractedDocument(metadata=metadata, sections=sections)


class DocxExtractor(DocumentExtractor):
    def extract(self, file_path: str | Path, document_name: str | None = None) -> ExtractedDocument:
        path = Path(file_path)
        filename = document_name or path.name
        document = Document(str(path))
        sections: list[ExtractedSection] = []
        buffer: list[str] = []
        current_title: str | None = None
        hierarchy: list[str] = []

        def flush() -> None:
            if not buffer:
                return
            sections.append(
                ExtractedSection(
                    content="\n".join(buffer),
                    document_name=filename,
                    section_title=current_title,
                    hierarchy=tuple(hierarchy),
                    metadata={"extraction_backend": "python-docx"},
                )
            )
            buffer.clear()

        for paragraph in document.paragraphs:
            text = " ".join(paragraph.text.split())
            if not text:
                continue

            heading_level = _docx_heading_level(paragraph.style.name if paragraph.style else "")
            if heading_level is not None:
                flush()
                hierarchy = hierarchy[: heading_level - 1]
                hierarchy.append(text)
                current_title = text
                continue

            buffer.append(text)

        flush()
        sections.extend(_docx_table_sections(document, filename, hierarchy))

        metadata = DocumentMetadata(filename=filename, file_type=".docx", total_pages=None)
        logger.info("Extracted DOCX filename=%s sections=%s", filename, len(sections))
        return ExtractedDocument(metadata=metadata, sections=sections)


class PdfExtractor(DocumentExtractor):
    def extract(self, file_path: str | Path, document_name: str | None = None) -> ExtractedDocument:
        path = Path(file_path)
        filename = document_name or path.name
        table_sections_by_page = _extract_pdf_tables(path, filename)
        sections: list[ExtractedSection] = []

        with fitz.open(path) as document:
            total_pages = document.page_count
            for page_index, page in enumerate(document, start=1):
                sections.extend(
                    _sections_from_pdf_page(
                        page=page,
                        document_name=filename,
                        page_number=page_index,
                        table_detected=bool(table_sections_by_page.get(page_index)),
                    )
                )
                sections.extend(table_sections_by_page.get(page_index, []))

        metadata = DocumentMetadata(filename=filename, file_type=".pdf", total_pages=total_pages)
        logger.info(
            "Extracted PDF filename=%s pages=%s sections=%s",
            filename,
            total_pages,
            len(sections),
        )
        return ExtractedDocument(metadata=metadata, sections=sections)


def get_extractor(file_path: str | Path) -> DocumentExtractor:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".txt":
        return TxtExtractor()
    if suffix == ".docx":
        return DocxExtractor()
    if suffix == ".pdf":
        return PdfExtractor()
    raise ValueError(f"Unsupported file format: {suffix}")


def _sections_from_text_lines(
    lines: list[str],
    document_name: str,
    page_number: int | None,
    metadata: dict[str, Any],
) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    buffer: list[str] = []
    current_title: str | None = None
    hierarchy: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        sections.append(
            ExtractedSection(
                content="\n".join(buffer),
                document_name=document_name,
                page_number=page_number,
                section_title=current_title,
                hierarchy=tuple(hierarchy),
                metadata=dict(metadata),
            )
        )
        buffer.clear()

    for raw_line in lines:
        line = " ".join(raw_line.split())
        if not line:
            continue
        if _looks_like_heading(line):
            flush()
            current_title = line
            hierarchy = [line]
            continue
        buffer.append(line)

    flush()
    if not sections and buffer:
        flush()
    return sections


def _sections_from_pdf_page(
    page,
    document_name: str,
    page_number: int,
    table_detected: bool,
) -> list[ExtractedSection]:
    line_items = _pdf_line_items(page)
    if not line_items:
        return []

    body_size = median([item["font_size"] for item in line_items])
    sections: list[ExtractedSection] = []
    buffer: list[str] = []
    current_title: str | None = None
    hierarchy: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        sections.append(
            ExtractedSection(
                content="\n".join(buffer),
                document_name=document_name,
                page_number=page_number,
                section_title=current_title,
                table_detected=table_detected,
                hierarchy=tuple(hierarchy),
                metadata={"extraction_backend": "pymupdf"},
            )
        )
        buffer.clear()

    for item in line_items:
        text = item["text"]
        if _looks_like_pdf_heading(text, item["font_size"], body_size):
            flush()
            current_title = text
            hierarchy = [text]
            continue
        buffer.append(text)

    flush()
    return sections


def _pdf_line_items(page) -> list[dict[str, Any]]:
    page_dict = page.get_text("dict")
    lines: list[dict[str, Any]] = []

    for block in page_dict.get("blocks", []):
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = " ".join(span.get("text", "").strip() for span in spans).strip()
            text = " ".join(text.split())
            if not text:
                continue
            font_sizes = [float(span.get("size", 0)) for span in spans if span.get("text", "").strip()]
            lines.append({"text": text, "font_size": max(font_sizes or [0.0])})

    return lines


def _extract_pdf_tables(path: Path, document_name: str) -> dict[int, list[ExtractedSection]]:
    try:
        import pdfplumber
    except ImportError:
        logger.warning("pdfplumber is not installed; PDF table extraction will be skipped.")
        return {}

    tables_by_page: dict[int, list[ExtractedSection]] = {}
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_tables = page.extract_tables() or []
            for table_index, table in enumerate(page_tables, start=1):
                rows = [
                    " | ".join("" if cell is None else str(cell).strip() for cell in row)
                    for row in table
                    if any(cell for cell in row)
                ]
                table_text = "\n".join(row for row in rows if row.strip())
                if not table_text:
                    continue
                tables_by_page.setdefault(page_number, []).append(
                    ExtractedSection(
                        content=table_text,
                        document_name=document_name,
                        page_number=page_number,
                        section_title=f"Table {table_index}",
                        table_detected=True,
                        hierarchy=(f"Table {table_index}",),
                        metadata={
                            "extraction_backend": "pdfplumber",
                            "table_index": table_index,
                        },
                    )
                )

    return tables_by_page


def _docx_table_sections(document: Document, filename: str, hierarchy: list[str]) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        table_text = "\n".join(rows)
        if not table_text:
            continue
        sections.append(
            ExtractedSection(
                content=table_text,
                document_name=filename,
                section_title=f"Table {table_index}",
                table_detected=True,
                hierarchy=tuple([*hierarchy, f"Table {table_index}"]),
                metadata={"extraction_backend": "python-docx", "table_index": table_index},
            )
        )
    return sections


def _docx_heading_level(style_name: str) -> int | None:
    match = re.match(r"Heading\s+(\d+)", style_name or "", flags=re.IGNORECASE)
    if not match:
        return None
    return int(match.group(1))


def _looks_like_heading(text: str) -> bool:
    if len(text) > 120:
        return False
    if re.match(r"^(\d+(\.\d+)*\.?|[A-Z]\.)\s+\S+", text):
        return True
    if text.isupper() and len(text.split()) <= 12:
        return True
    if text.istitle() and not text.endswith((".", ",", ";", ":")) and len(text.split()) <= 10:
        return True
    return False


def _looks_like_pdf_heading(text: str, font_size: float, body_size: float) -> bool:
    if len(text) > 140:
        return False
    if font_size >= body_size * 1.18 and len(text.split()) <= 14:
        return True
    return _looks_like_heading(text)
